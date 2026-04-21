#!/usr/bin/env python3
"""Extract raw text and metadata from prompt or local project-plan files."""

from __future__ import annotations

import argparse
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from repo_init_common import (
    VERSION,
    first_non_empty_line,
    normalize_text,
    read_text_with_fallback,
    unique,
    utc_now_iso,
    write_json,
)


class SimpleHtmlTextExtractor(HTMLParser):
    """Extract title, headings, and text from HTML."""

    BLOCK_TAGS = {
        "article",
        "body",
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "li",
        "main",
        "p",
        "section",
        "tr",
    }
    HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self) -> None:
        super().__init__()
        self._skip_depth = 0
        self._title_depth = 0
        self._title_parts: list[str] = []
        self._heading_tag: str | None = None
        self._heading_parts: list[str] = []
        self._text_parts: list[str] = []
        self.headings: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip_depth += 1
            return
        if tag == "title":
            self._title_depth += 1
        if tag in self.HEADING_TAGS:
            self._heading_tag = tag
            self._heading_parts = []
        if tag in self.BLOCK_TAGS:
            self._text_parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip_depth:
            self._skip_depth -= 1
            return
        if tag == "title" and self._title_depth:
            self._title_depth -= 1
        if tag == self._heading_tag:
            heading = normalize_text("".join(self._heading_parts))
            if heading:
                self.headings.append(heading)
                self._text_parts.append(heading)
                self._text_parts.append("\n")
            self._heading_tag = None
            self._heading_parts = []

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        if self._title_depth:
            self._title_parts.append(data)
            return
        if self._heading_tag:
            self._heading_parts.append(data)
            return
        self._text_parts.append(data)

    @property
    def title(self) -> str:
        return normalize_text("".join(self._title_parts))

    @property
    def text(self) -> str:
        return normalize_text("".join(self._text_parts))


def extract_prompt(prompt: str) -> dict[str, Any]:
    """Extract intake information from prompt-only input."""
    raw_text = normalize_text(prompt)
    title = first_non_empty_line(raw_text)[:120] or "Prompt input"
    return {
        "source_type": "prompt",
        "source_path": None,
        "detected_format": "text",
        "title": title,
        "raw_text": raw_text,
        "warnings": [],
        "metadata": {
            "headings": [],
        },
        "provenance": {
            "extraction_timestamp": utc_now_iso(),
            "extractor": "extract_project_source.py",
            "extractor_version": VERSION,
            "known_limitations": [
                "Prompt-only intake relies on natural-language structure supplied by the user.",
            ],
        },
    }


def extract_text_file(path: Path, detected_format: str) -> dict[str, Any]:
    """Extract plain text or markdown."""
    text, warnings = read_text_with_fallback(path)
    raw_text = normalize_text(text)
    return {
        "source_type": "file",
        "source_path": str(path.resolve()),
        "detected_format": detected_format,
        "title": first_non_empty_line(raw_text)[:120] or path.stem,
        "raw_text": raw_text,
        "warnings": unique(warnings),
        "metadata": {
            "headings": [],
            "file_size_bytes": path.stat().st_size,
        },
        "provenance": {
            "extraction_timestamp": utc_now_iso(),
            "extractor": "extract_project_source.py",
            "extractor_version": VERSION,
            "known_limitations": [
                "Plain text and markdown extraction preserves text but not semantic relationships beyond line structure.",
            ],
        },
    }


def extract_html_file(path: Path) -> dict[str, Any]:
    """Extract text from HTML."""
    text, warnings = read_text_with_fallback(path)
    parser = SimpleHtmlTextExtractor()
    parser.feed(text)
    raw_text = parser.text
    if not raw_text:
        warnings.append(f"No visible text was extracted from {path.name}.")
    return {
        "source_type": "file",
        "source_path": str(path.resolve()),
        "detected_format": "html",
        "title": parser.title or first_non_empty_line(raw_text)[:120] or path.stem,
        "raw_text": raw_text,
        "warnings": unique(warnings),
        "metadata": {
            "headings": parser.headings,
            "file_size_bytes": path.stat().st_size,
        },
        "provenance": {
            "extraction_timestamp": utc_now_iso(),
            "extractor": "extract_project_source.py",
            "extractor_version": VERSION,
            "known_limitations": [
                "HTML extraction strips layout and may flatten nested content.",
            ],
        },
    }


def extract_docx_file(path: Path) -> dict[str, Any]:
    """Extract text from DOCX files."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("python-docx is required to extract DOCX files.") from exc

    document = Document(str(path))
    paragraphs: list[str] = []
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        paragraphs.append(text)
        style_name = getattr(paragraph.style, "name", "")
        if isinstance(style_name, str) and style_name.lower().startswith("heading"):
            headings.append(text)

    table_chunks: list[str] = []
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            table_chunks.append("\n".join(rows))

    raw_text = normalize_text("\n\n".join(paragraphs + table_chunks))
    warnings: list[str] = []
    if table_chunks:
        warnings.append("DOCX tables were flattened into plain text rows.")
    if not raw_text:
        warnings.append(f"No text was extracted from {path.name}.")

    core_props = document.core_properties
    title = core_props.title or (headings[0] if headings else None)
    title = title or first_non_empty_line(raw_text)[:120] or path.stem
    return {
        "source_type": "file",
        "source_path": str(path.resolve()),
        "detected_format": "docx",
        "title": title,
        "raw_text": raw_text,
        "warnings": unique(warnings),
        "metadata": {
            "headings": headings,
            "paragraph_count": len(paragraphs),
            "file_size_bytes": path.stat().st_size,
        },
        "provenance": {
            "extraction_timestamp": utc_now_iso(),
            "extractor": "extract_project_source.py",
            "extractor_version": VERSION,
            "known_limitations": [
                "DOCX extraction preserves paragraph text but does not retain tracked formatting semantics.",
            ],
        },
    }


def extract_pdf_file(path: Path) -> dict[str, Any]:
    """Extract text from PDF files."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("pypdf is required to extract PDF files.") from exc

    reader = PdfReader(str(path))
    warnings: list[str] = []
    if reader.is_encrypted:
        decrypt_status = reader.decrypt("")
        warnings.append("PDF is encrypted; attempted empty-password decryption.")
        if decrypt_status == 0:
            warnings.append("PDF could not be decrypted with an empty password.")

    page_texts: list[str] = []
    empty_pages = 0
    for index, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            page_texts.append(text)
        else:
            empty_pages += 1
            warnings.append(f"No extractable text detected on PDF page {index}.")

    raw_text = normalize_text("\n\n".join(page_texts))
    metadata = reader.metadata or {}
    title = (getattr(metadata, "title", None) or metadata.get("/Title")) if metadata else None
    if isinstance(title, str) and title.strip().lower() == "untitled":
        title = None
    title = title or first_non_empty_line(raw_text)[:120] or path.stem
    if not raw_text:
        warnings.append(f"PDF extraction produced no usable text for {path.name}.")
    if empty_pages:
        warnings.append("PDF may require OCR or manual review for some pages.")

    return {
        "source_type": "file",
        "source_path": str(path.resolve()),
        "detected_format": "pdf",
        "title": title,
        "raw_text": raw_text,
        "warnings": unique(warnings),
        "metadata": {
            "headings": [],
            "page_count": len(reader.pages),
            "file_size_bytes": path.stat().st_size,
        },
        "provenance": {
            "extraction_timestamp": utc_now_iso(),
            "extractor": "extract_project_source.py",
            "extractor_version": VERSION,
            "known_limitations": [
                "PDF extraction may lose layout, tables, and scanned text without OCR.",
            ],
        },
    }


def extract_from_path(path: Path) -> dict[str, Any]:
    """Dispatch extraction by file type."""
    if not path.exists():
        raise FileNotFoundError(f"Input path does not exist: {path}")
    if not path.is_file():
        raise ValueError(f"Input path is not a file: {path}")

    suffix = path.suffix.lower()
    if suffix == ".md":
        return extract_text_file(path, "markdown")
    if suffix == ".txt":
        return extract_text_file(path, "text")
    if suffix == ".html" or suffix == ".htm":
        return extract_html_file(path)
    if suffix == ".docx":
        return extract_docx_file(path)
    if suffix == ".pdf":
        return extract_pdf_file(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--prompt", help="Natural-language project initialization prompt.")
    source.add_argument("--path", help="Path to a local project-plan file.")
    parser.add_argument("--output", help="Optional path to write JSON output.")
    return parser.parse_args()


def main() -> int:
    """Run the extractor."""
    args = parse_args()
    try:
        result = extract_prompt(args.prompt) if args.prompt else extract_from_path(Path(args.path))
        write_json(result, args.output)
        return 0
    except Exception as exc:  # pragma: no cover - CLI entry point
        print(f"extract_project_source.py: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
